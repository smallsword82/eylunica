#!/usr/bin/env python3
"""Build grade-5 topics from official weekly plans and grade-filtered programs."""
import io,json,pathlib,re,hashlib,zipfile,unicodedata
import openpyxl,xlrd
ROOT=pathlib.Path(__file__).resolve().parents[2];CACHE=ROOT/'.tools/curriculum-sources'
s=json.loads((ROOT/'src/grade5/official-curriculum.json').read_text())
raw=CACHE/'raw-curriculum.json'
if not raw.exists():raw.write_text(json.dumps(s,ensure_ascii=False))
else:s=json.loads(raw.read_text())
map_ids={'ortaokul-matematik':'course-1','ortaokul-turkce':'course-0','fen-bilimleri':'course-2','sosyal-bilgiler':'course-3','coklu-yabanci-dil-egitim-modeli-almanca':'course-5','coklu-yabanci-dil-egitim-modeli-ingilizce':'meb-course-35','ingilizce-dersi-yillik':'course-4','beden-egitimi-ve-spor':'course-10','bilisim-teknolojileri':'course-7','gorsel-sanatlar':'course-8','muzik-dersi-yillik':'course-9'}
def clean(v):return ' '.join(str(v or '').split())
def norm(v):return clean(v).casefold().replace('ı','i').replace('i̇','i')
def titles(content,cid):
 if cid in ['course-4','meb-course-35']:
  m=re.search(r'Sub[ -]?themes?\s*:\s*(.*?)(?:Target Vocabulary|Vocabulary|Target Language|\n[A-Z ]{10,}:)',content,re.I|re.S)
  if m:content=clean(m.group(1)).strip(' “”.')
  else:return []
 elif cid=='course-5':
  m=re.search(r'Unterthemen\s*:\s*(.*?)(?:Zielwortschatz|\n\w+\s*im Gebrauch)',content,re.I|re.S)
  if m:content=clean(m.group(1))
  else:return []
 content=re.sub(r'\s*\*\s*', '\n', content)
 return list(dict.fromkeys(x.strip(' *•–-.:“”') for x in re.split(r'[;\n]',content) if 3<len(x.strip(' *•–-.:“”'))<200))
def outputs(text):
 return [dict(code=m.group(1).rstrip('.'),text=m.group(2).strip()) for m in re.finditer(r'(?m)^\s*([A-ZÇĞİÖŞÜ]{1,6}(?:\.[A-ZÇĞİÖŞÜ])?\.5\.[A-Za-z\d.]+)[.\s]+([^\n]+)',text)]
def identifier(cid,unit,title):return 'topic-'+hashlib.sha256((cid+'|'+norm(unit)+'|'+norm(title)).encode()).hexdigest()[:16]
records={};plan_courses=set();sources=[]
for p in sorted(CACHE.glob('*.zip')):
 cid=next((v for k,v in map_ids.items() if k in p.name),None)
 if not cid:continue
 source='https://tymm.meb.gov.tr/assets/file/'+p.name
 for filename in zipfile.ZipFile(p).namelist():
  if not filename.endswith(('.xlsx','.xls')) or 'Hecce' in filename:continue
  data=zipfile.ZipFile(p).read(filename)
  if filename.endswith('.xlsx'):
   book=openpyxl.load_workbook(io.BytesIO(data),data_only=True)
   sheets=[]
   for w in book:
    if not re.search(r'(^|[^0-9])5([^0-9]|$)',w.title):continue
    vals=[[w.cell(r,c).value for c in range(1,min(w.max_column,25)+1)] for r in range(1,min(w.max_row,1000)+1)]
    for area in w.merged_cells.ranges:
     if area.min_col>12:continue
     v=w.cell(area.min_row,area.min_col).value
     for r in range(area.min_row,min(area.max_row,len(vals))+1):
      for c in range(area.min_col,min(area.max_col,len(vals[r-1]))+1):vals[r-1][c-1]=v
    sheets.append((w.title,vals))
  else:
   book=xlrd.open_workbook(file_contents=data)
   sheets=[(w.name,[w.row_values(r) for r in range(w.nrows)]) for w in book.sheets() if re.search(r'(^|[^0-9])5([^0-9]|$)',w.name)]
  for sheet,rows in sheets:
   plan_courses.add(cid);sources.append(dict(courseId=cid,url=source,file=filename,sheet=sheet,sha256=hashlib.sha256(data).hexdigest()))
   for rownum,row in enumerate(rows,1):
    if len(row)<6:continue
    m=re.search(r'(\d+)\.?\s*(?:Hafta|Woche)',str(row[1]),re.I) or re.search(r'Week\s*(\d+)',str(row[1]),re.I)
    if not m:continue
    week=int(m.group(1))
    if not 1<=week<=36:continue
    unit=clean(row[3]);content=str(row[5] or '') if cid=='course-0' else str(row[4] or '')
    if not unit or any(x in norm(unit) for x in ['orientation','okul temelli','tatil','uyum','wiederholung']):continue
    names=titles(content,cid)
    outcomeText='\n'.join(str(v or '') for v in row[6:10]) if cid=='course-0' else str(row[5] or '')
    for title in names:
     key=identifier(cid,unit,title)
     if key not in records:records[key]=dict(id=key,courseId=cid,unit=unit,title=title,year='2026-2027',models=[next(c['model'] for c in s['courses'] if c['id']==cid)],weeks=[],outcomes=[],source=dict(url=source,checkedAt=s['checkedAt'],status='verified',note='MEB 2026–2027 taslak yıllık planı; okul/zümre uyarlamaları olabilir.',sheet=sheet,rows=[]),summary='',example='',outcomeScope='week',dates=[])
     r=records[key];r['weeks']=sorted(set(r['weeks']+[week]));r['source']['rows'].append(rownum)
     date=clean(row[1])
     if date not in r['dates']:r['dates'].append(date)
     for out in outputs(outcomeText):
      if out['code'] not in [x['code'] for x in r['outcomes']]:r['outcomes'].append(out)
# For programs without a published plan, preserve exact content-frame topics.
# No invented weeks. Language vocabulary lists are not treated as lesson topics.
for u in s['units']:
 if u['courseId'] in plan_courses:continue
 content=u['framework'];names=titles(content,u['courseId'])
 if u['courseId'].startswith('meb-course-11'):
  m=re.search(r'Alt Temalar\s*:?\s*(.*?)(?:İşlevsel|Hedeflenen|Hedef Dil|Hedef Kelime|Hedef Sözcük|$)', content, re.S)
  if m:names=[l.strip('0123456789. ') for l in m.group(1).splitlines() if 2<len(l.strip())<180]
 if not names:names=[u['title']]
 for title in names:
  key=identifier(u['courseId'],u['unit'],title)
  records[key]={**{k:v for k,v in u.items() if k in ['courseId','unit','year','models','source','summary','outcomes','weeks']},'id':key,'title':title,'example':'','outcomeScope':'unit'}
# Compact unit index retains every official unit and code without long repeated fields.
units=[{k:v for k,v in u.items() if k not in ['outcomeText','framework','subtopics']} for u in s['units']]
s['units']=units;s['topics']=list(records.values());s['annualPlans']=sources
board=json.loads((CACHE/'timetable-source.json').read_text())
s['timetable']={**board,'decision':'115','decisionDate':'2026-09-23','checkedAt':'2026-10-21','pdfSha256':hashlib.sha256((CACHE/'timetable-2026.pdf').read_bytes()).hexdigest(),'mandatoryTotalHours':30,'electiveTotalHours':5,'totalHours':35}
hours={'course-0':6,'course-1':5,'course-2':4,'course-3':3,'course-4':3,'course-6':2,'course-7':2,'course-8':1,'course-9':1,'course-10':2}
for course in s['courses']:
 course['weeklyHours']=hours.get(course['id']);course['category']='mandatory' if course['id'] in hours else 'multiple-language' if course['model']=='multiple-language' else 'elective'
 course['timetableSource']=board['url']
s['courses'].append(dict(id='guidance',name='Rehberlik ve Yönlendirme',model='standard',verification='verified',weeklyHours=1,category='mandatory',sourceUrl=board['url'],programUrl=board['url'],checkedAt='2026-10-21'))

s['scheduleVerification']='verified-2026-2027-decision-115'
(ROOT/'src/grade5/official-curriculum.json').write_text(json.dumps(s,ensure_ascii=False,indent=2)+'\n')
print(json.dumps(dict(courses=len(s['courses']),units=len(units),topics=len(records),withWeeks=sum(bool(r['weeks']) for r in records.values()),outcomes=sum(len(u['outcomes']) for u in units)),ensure_ascii=False))
for c in s['courses']:print(c['name'],len([r for r in records.values() if r['courseId']==c['id']]))
